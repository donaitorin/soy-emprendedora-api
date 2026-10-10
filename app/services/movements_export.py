"""Builds the .xlsx for GET /accounts/{id}/movements/export."""

from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from app.models.money_movement import MoneyMovement, MovementType

GREEN = "15803D"
RED = "B91C1C"
CURRENCY_FORMAT = '"$"#,##0.00'
DATE_FORMAT = "yyyy-mm-dd"

HEADERS = ["Fecha", "Tipo", "Monto", "Fuente / Categoría", "Método de pago"]

_LABELS = {
    "mentoria": "Mentoría",
    "comunidad": "Comunidad",
    "claridad": "Claridad",
    "producto": "Producto",
    "otro": "Otro",
    "transferencia": "Transferencia",
    "stripe": "Stripe",
    "mercadopago": "MercadoPago",
    "paypal": "PayPal",
    "efectivo": "Efectivo",
    "herramientas": "Herramientas",
    "publicidad": "Publicidad",
    "educacion": "Educación",
    "servicios": "Servicios",
}


def _label(value) -> str:
    if value is None:
        return ""
    return _LABELS.get(value.value, value.value)


def build_movements_workbook(movements: list[MoneyMovement]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Movimientos"

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    total_income = Decimal(0)
    total_expense = Decimal(0)

    for m in movements:
        is_income = m.type == MovementType.INCOME
        amount = Decimal(m.amount)
        if is_income:
            total_income += amount
        else:
            total_expense += amount

        ws.append(
            [
                m.occurred_on,
                "Ingreso" if is_income else "Gasto",
                float(amount),
                _label(m.source if is_income else m.category),
                _label(m.payment_method) if is_income else "",
            ]
        )
        row = ws.max_row
        ws.cell(row, 1).number_format = DATE_FORMAT
        amount_cell = ws.cell(row, 3)
        amount_cell.number_format = CURRENCY_FORMAT
        amount_cell.font = Font(color=GREEN if is_income else RED)

    ws.append([])
    for label, value, color, bold in (
        ("Total ingresos", total_income, GREEN, False),
        ("Total gastos", total_expense, RED, False),
        ("Resultado", total_income - total_expense, None, True),
    ):
        ws.append([label, None, float(value)])
        row = ws.max_row
        ws.cell(row, 1).font = Font(bold=True)
        cell = ws.cell(row, 3)
        cell.number_format = CURRENCY_FORMAT
        cell.font = Font(bold=bold, color=color)

    for col, width in zip("ABCDE", (14, 10, 16, 22, 18), strict=True):
        ws.column_dimensions[col].width = width

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
